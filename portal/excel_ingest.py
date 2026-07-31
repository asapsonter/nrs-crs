"""Excel and CSV ingest for CRS filings.

Follows the bulk-template pattern used by AEOI portals such as the Cayman
DITC: a fixed header row, one account per data row, with optional controlling
person columns for Passive NFE (CRS101) accounts. Accepts .xlsx workbooks and
the same template saved as .csv.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.utils.dateparse import parse_date

from portal.models import AccountReport

# Template columns. Keys are normalised (lower, no spaces/underscores).
REQUIRED_COLUMNS = [
    "HolderName",
    "ResidenceCountry",
    "Address",
    "AccountNumber",
    "Currency",
]
OPTIONAL_COLUMNS = [
    "HolderType",
    "AcctHolderType",
    # NamePerson_Type components. Supplying these lets the record be reported
    # with a real FirstName/LastName pair instead of the NFN fallback.
    "FirstName",
    "LastName",
    "TIN",
    "TINMissingReason",
    "AddressCountry",
    # AddressFix components — City is what promotes a record from AddressFree.
    "Street",
    "BuildingIdentifier",
    "PostCode",
    "City",
    "CountrySubentity",
    "BirthDate",
    "BirthCity",
    "BirthCitySubentity",
    "BirthCountry",
    "BirthFormerCountryName",
    "SelfCertification",
    "AcctNumberType",
    "ClosedAccount",
    "DormantAccount",
    "Balance",
    "Dividends",
    "Interest",
    "GrossProceeds",
    "OtherIncome",
    "CPName",
    "CPFirstName",
    "CPLastName",
    "CPResidence",
    "CPTIN",
    "CPCity",
    "CPType",
]
ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS

_HOLDER_TYPES = {"INDIVIDUAL", "ORGANISATION"}
_ACCT_HOLDER_TYPES = {"CRS101", "CRS102", "CRS103"}
_SELF_CERTS = {"OBTAINED", "CURED", "NOT_OBTAINED"}
_ACCT_NUMBER_TYPES = set(AccountReport.AcctNumberType.values)
_TRUE_VALUES = {"TRUE", "YES", "Y", "1"}
_FALSE_VALUES = {"", "FALSE", "NO", "N", "0"}


def _norm(name: str) -> str:
    return "".join(ch for ch in str(name or "").lower() if ch.isalnum())


@dataclass
class ParsedExcelRecord:
    holder_name: str = ""
    holder_first_name: str = ""
    holder_last_name: str = ""
    holder_type: str = "INDIVIDUAL"
    acct_holder_type: str = ""
    residence_country: str = ""
    foreign_tin: str = ""
    tin_unavailable_reason: str = ""
    holder_address: str = ""
    address_country: str = ""
    holder_street: str = ""
    holder_building_identifier: str = ""
    holder_post_code: str = ""
    holder_city: str = ""
    holder_country_subentity: str = ""
    birth_date: object = None
    birth_city: str = ""
    birth_city_subentity: str = ""
    birth_country_code: str = ""
    birth_former_country_name: str = ""
    self_certification: str = ""
    account_number: str = ""
    acct_number_type: str = ""
    closed_account: bool = False
    dormant_account: bool = False
    currency: str = ""
    balance: Decimal = Decimal("0")
    dividends: Decimal = Decimal("0")
    interest: Decimal = Decimal("0")
    gross_proceeds: Decimal = Decimal("0")
    other_income: Decimal = Decimal("0")
    cp_name: str = ""
    cp_first_name: str = ""
    cp_last_name: str = ""
    cp_residence: str = ""
    cp_tin: str = ""
    cp_city: str = ""
    cp_type: str = ""


@dataclass
class ExcelParseResult:
    ok: bool
    errors: list[str] = field(default_factory=list)
    records: list[ParsedExcelRecord] = field(default_factory=list)


def _rows_from_csv(content: bytes) -> tuple[list[list[str]], str]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return [], "The CSV file is not valid UTF-8 text."
    reader = csv.reader(io.StringIO(text))
    return [row for row in reader if any(str(cell).strip() for cell in row)], ""


def _rows_from_xlsx(content: bytes) -> tuple[list[list[str]], str]:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return [], "The file could not be read as an Excel workbook (.xlsx)."
    sheet = workbook.worksheets[0]
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        cells = ["" if cell is None else str(cell).strip() for cell in row]
        if any(cells):
            rows.append(cells)
    workbook.close()
    return rows, ""


def parse_excel_upload(content: bytes, filename: str) -> ExcelParseResult:
    """Parse and validate an uploaded Excel or CSV CRS filing template."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        rows, error = _rows_from_csv(content)
    elif name.endswith(".xlsx"):
        rows, error = _rows_from_xlsx(content)
    else:
        return ExcelParseResult(ok=False, errors=["Upload a .xlsx workbook or a .csv file using the CRS template."])
    if error:
        return ExcelParseResult(ok=False, errors=[error])
    if not rows:
        return ExcelParseResult(ok=False, errors=["The file is empty."])

    header = [_norm(cell) for cell in rows[0]]
    col_index: dict[str, int] = {}
    for column in ALL_COLUMNS:
        key = _norm(column)
        if key in header:
            col_index[column] = header.index(key)

    errors: list[str] = []
    missing = [column for column in REQUIRED_COLUMNS if column not in col_index]
    if missing:
        errors.append(
            "Missing required column(s): " + ", ".join(missing)
            + ". The header row must use the CRS template column names."
        )
        return ExcelParseResult(ok=False, errors=errors)

    data_rows = rows[1:]
    if not data_rows:
        return ExcelParseResult(ok=False, errors=["The file has a header row but no account rows. Use a nil return if there is nothing to report."])

    def cell(row: list[str], column: str) -> str:
        index = col_index.get(column)
        if index is None or index >= len(row):
            return ""
        return str(row[index]).strip()

    records: list[ParsedExcelRecord] = []
    for line, row in enumerate(data_rows, start=2):
        record = ParsedExcelRecord()
        record.holder_name = cell(row, "HolderName")
        record.residence_country = cell(row, "ResidenceCountry").upper()
        record.holder_address = cell(row, "Address")
        record.account_number = cell(row, "AccountNumber")
        record.currency = cell(row, "Currency").upper()
        for label, value in (
            ("HolderName", record.holder_name),
            ("ResidenceCountry", record.residence_country),
            ("Address", record.holder_address),
            ("AccountNumber", record.account_number),
            ("Currency", record.currency),
        ):
            if not value:
                errors.append(f"Row {line}: {label} is required.")
        if record.residence_country and len(record.residence_country) != 2:
            errors.append(f"Row {line}: ResidenceCountry must be a 2-letter ISO code.")

        holder_type = cell(row, "HolderType").upper() or "INDIVIDUAL"
        if holder_type not in _HOLDER_TYPES:
            errors.append(f"Row {line}: HolderType must be INDIVIDUAL or ORGANISATION.")
            holder_type = "INDIVIDUAL"
        record.holder_type = holder_type

        acct_holder_type = cell(row, "AcctHolderType").upper()
        if acct_holder_type and acct_holder_type not in _ACCT_HOLDER_TYPES:
            errors.append(f"Row {line}: AcctHolderType must be CRS101, CRS102 or CRS103.")
            acct_holder_type = ""
        # Account Holder Type applies to entity holders only; it must be left
        # blank for individuals (per the CRS portal convention).
        record.acct_holder_type = acct_holder_type if holder_type == "ORGANISATION" else ""

        record.holder_first_name = cell(row, "FirstName")
        record.holder_last_name = cell(row, "LastName")
        record.foreign_tin = cell(row, "TIN")
        record.tin_unavailable_reason = cell(row, "TINMissingReason")
        record.address_country = cell(row, "AddressCountry").upper() or record.residence_country

        record.holder_street = cell(row, "Street")
        record.holder_building_identifier = cell(row, "BuildingIdentifier")
        record.holder_post_code = cell(row, "PostCode")
        record.holder_city = cell(row, "City")
        record.holder_country_subentity = cell(row, "CountrySubentity")

        raw_birth = cell(row, "BirthDate")
        if raw_birth:
            parsed = parse_date(raw_birth[:10])
            if parsed is None:
                errors.append(f"Row {line}: BirthDate '{raw_birth}' is not a valid date (use YYYY-MM-DD).")
            else:
                record.birth_date = parsed
        record.birth_city = cell(row, "BirthCity")
        record.birth_city_subentity = cell(row, "BirthCitySubentity")
        record.birth_country_code = cell(row, "BirthCountry").upper()
        record.birth_former_country_name = cell(row, "BirthFormerCountryName")
        if record.birth_country_code and len(record.birth_country_code) != 2:
            errors.append(f"Row {line}: BirthCountry must be a 2-letter ISO code.")
            record.birth_country_code = ""
        if record.birth_country_code and record.birth_former_country_name:
            errors.append(
                f"Row {line}: give either BirthCountry or BirthFormerCountryName, not both "
                "— CRS CountryInfo is a choice between the two."
            )

        self_cert = cell(row, "SelfCertification").upper()
        if self_cert and self_cert not in _SELF_CERTS:
            errors.append(f"Row {line}: SelfCertification must be OBTAINED, CURED or NOT_OBTAINED.")
            self_cert = ""
        record.self_certification = self_cert

        acct_number_type = cell(row, "AcctNumberType").upper()
        if acct_number_type and acct_number_type not in _ACCT_NUMBER_TYPES:
            errors.append(
                f"Row {line}: AcctNumberType must be one of "
                + "; ".join(label for _, label in AccountReport.AcctNumberType.choices)
                + "."
            )
            acct_number_type = ""
        record.acct_number_type = acct_number_type

        for column, attr in (("ClosedAccount", "closed_account"), ("DormantAccount", "dormant_account")):
            raw_flag = cell(row, column).upper()
            if raw_flag in _TRUE_VALUES:
                setattr(record, attr, True)
            elif raw_flag not in _FALSE_VALUES:
                errors.append(f"Row {line}: {column} must be TRUE or FALSE.")

        decimal_attrs = {
            "Balance": "balance",
            "Dividends": "dividends",
            "Interest": "interest",
            "GrossProceeds": "gross_proceeds",
            "OtherIncome": "other_income",
        }
        for column, attr in decimal_attrs.items():
            raw = cell(row, column) or "0"
            try:
                setattr(record, attr, Decimal(raw.replace(",", "")))
            except InvalidOperation:
                errors.append(f"Row {line}: {column} value '{raw}' is not a valid amount.")

        record.cp_name = cell(row, "CPName")
        record.cp_first_name = cell(row, "CPFirstName")
        record.cp_last_name = cell(row, "CPLastName")
        record.cp_residence = cell(row, "CPResidence").upper()
        record.cp_tin = cell(row, "CPTIN")
        record.cp_city = cell(row, "CPCity")
        record.cp_type = cell(row, "CPType").upper()
        # A CP given only as first/last name still counts as present. A first
        # name alone is an error rather than a silent drop: LastName is the
        # mandatory half of the CRS NamePerson pair.
        if not record.cp_name and record.cp_last_name:
            record.cp_name = " ".join(part for part in (record.cp_first_name, record.cp_last_name) if part)
        elif not record.cp_name and record.cp_first_name:
            errors.append(
                f"Row {line}: CPFirstName was given without CPLastName or CPName; "
                "add the controlling person's last name (or full CPName)."
            )
        if record.cp_name and not record.cp_residence:
            errors.append(f"Row {line}: CPResidence is required when a CPName is given.")
        if record.cp_name and record.acct_holder_type != "CRS101":
            errors.append(
                f"Row {line}: controlling persons are only reported for a Passive NFE "
                "with controlling persons (HolderType ORGANISATION, AcctHolderType CRS101)."
            )

        records.append(record)

    if errors:
        return ExcelParseResult(ok=False, errors=errors)
    return ExcelParseResult(ok=True, records=records)
